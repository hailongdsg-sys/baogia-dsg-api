#!/usr/bin/env python3
"""
build_bao_gia.py

Dien du lieu vao file mau BBG_mau.xlsm de tao bao gia noi that Hoa Phat / The One.
Dong goi toan bo logic da kiem chung (chen dong khi >3 san pham, dich merged cells
dung cach, cong thuc tong ket, resize/anchor anh san pham) de goi tu node
"Execute Command" trong n8n hoac tu API (main.py).

Cach dung:
    python3 build_bao_gia.py --config config.json

File config.json (vi du):
{
  "template_path": "/data/templates/BBG_mau.xlsm",
  "output_path": "/data/output/BG_ChiLan.xlsm",
  "customer_name": "Chị Lan",
  "phone": "0901234567",
  "company": "",
  "tax_address": "",
  "delivery_address": "123 Nguyễn Văn A, Q.1, TP.HCM",
  "mst": "",
  "email": "",
  "shipping_fee": 0,
  "discount_percent": 0,
  "products": [
    {
      "code": "X100",
      "description": "Mo ta san pham lay tu hoaphatsaigon.com",
      "color": "",
      "qty": 5,
      "unit_price": 1500000,
      "image_path": "/data/images/x100.png"
    }
  ]
}

Ghi chu:
- Neu san pham co "image_url" thay vi "image_path", script se tu tai anh ve
  (dung requests - server chay n8n thuong khong bi chan mang nhu sandbox Claude,
  nen KHONG can chup man hinh + crop nhu quy trinh thu cong).
- Neu KHONG co anh (ca image_path lan image_url deu thieu, hoac tai anh loi),
  script se BO QUA anh cho dong san pham do va van tiep tuc tao bao gia binh
  thuong (KHONG dung sys.exit/crash toan bo tien trinh nua - fix quan trong,
  vi truoc day 1 san pham thieu anh se lam sap toan bo API server).
- discount_percent (vi du 2 = giam 2%) se duoc dien vao o "Chiet khau" duoi
  dang cong thuc =Tong_cong*X%, tu tinh lai neu tong cong thay doi.
- KHONG duoc sua doi cong thuc tren cac o khac ngoai nhung o duoc liet ke trong
  script nay - day la file .xlsm co san cong thuc VAT/tam ung, sua sai se lam
  bao gia tinh sai tien.
"""
import argparse
import json
import os
import re
import sys
import zipfile
from copy import copy

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.utils import range_boundaries, get_column_letter, column_index_from_string

SHEET_NAME = "BG mau"
FIRST_PRODUCT_ROW = 18
DEFAULT_TEMPLATE_ROWS = 3          # so dong san pham co san trong template (18,19,20)
DEFAULT_TOTAL_START_ROW = 21       # dong "Tong cong" mac dinh khi khong chen them dong
IMG_WIDTH = 140
IMG_HEIGHT = 146
ROW_HEIGHT_WITH_IMAGE = 170

# --- Tu dong chinh chieu cao dong cho khu vuc thong tin khach hang (dong 10-14:
# Kinh gui/Cong ty/Dia chi thue/Dia chi giao hang/MST) - cac o nay wrap_text=True
# nhung template co san CHI CO CHIEU CAO 1 DONG CO DINH, nen khi noi dung dai
# (vd dia chi thue day du) se bi wrap xuong 2-3 dong nhung dong bi cat/de len
# nhau tren PDF. Uoc luong so dong can thiet bang do do rong chu that su (Pillow)
# roi nhan len theo BASE_ROW_HEIGHT.
BASE_ROW_HEIGHT = 15.0
INFO_ROW_FONT_SIZE = 11
_FONT_CANDIDATES = {
    "bold": [
        "/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSerif-Bold.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf",
    ],
    "regular": [
        "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSerif-Regular.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf",
    ],
    "italic": [
        "/usr/share/fonts/truetype/liberation/LiberationSerif-Italic.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSerif-Italic.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Italic.ttf",
    ],
}


def _load_font(style, size_px):
    """style: 'bold' | 'regular' | 'italic'."""
    from PIL import ImageFont
    for path in _FONT_CANDIDATES.get(style, []):
        try:
            return ImageFont.truetype(path, size_px)
        except Exception:
            continue
    return None


def _load_bold_font(size_px):
    return _load_font("bold", size_px)


def _merged_width_px(ws, cols):
    """Uoc luong be rong (px) cua vung merge gom cac cot trong `cols`, dung cung
    cong thuc quy doi "do rong Excel (ky tu)" -> pixel da dung cho anh san pham
    (width_chars * 7 + 5)."""
    total_chars = 0.0
    for col in cols:
        dim = ws.column_dimensions.get(col)
        total_chars += dim.width if dim and dim.width else 8.43
    return total_chars * 7 + 5


def _wrap_line_count(text, font, max_width_px):
    """Dem so dong can thiet cho 1 DOAN (khong chua "\\n") sau khi wrap theo tu."""
    if not text:
        return 1
    if font is None:
        return max(1, -(-len(text) // 90))
    words = text.split(" ")
    lines = 1
    cur = ""
    for w in words:
        trial = (cur + " " + w).strip()
        width = font.getlength(trial) if hasattr(font, "getlength") else font.getsize(trial)[0]
        if cur and width > max_width_px:
            lines += 1
            cur = w
        else:
            cur = trial
    return lines


def _count_multiline_wrapped(text, font, max_width_px):
    """Nhu _wrap_line_count nhung TON TRONG cac dau xuong dong "\\n" co san trong
    text (moi doan giua 2 dau \\n duoc wrap rieng, dong trong tinh la 1 dong)."""
    if not text:
        return 0
    total = 0
    for seg in text.split("\n"):
        if seg.strip() == "":
            total += 1
        else:
            total += _wrap_line_count(seg, font, max_width_px)
    return total


def _count_wrapped_lines(text, font_size_pt, max_width_px):
    font = _load_bold_font(round(font_size_pt * 96 / 72))
    return max(1, _count_multiline_wrapped(text, font, max_width_px))


def set_wrapped_row_height(ws, row, text, cols="ABCDEF", font_size_pt=INFO_ROW_FONT_SIZE):
    """Dat lai chieu cao dong `row` sao cho du hien thi het `text` sau khi wrap
    trong vung merge `cols`, tranh bi cat/de chong len dong ben duoi."""
    max_width_px = max(20, _merged_width_px(ws, cols) - 10)  # tru le trong o
    n_lines = _count_wrapped_lines(text, font_size_pt, max_width_px)
    ws.row_dimensions[row].height = BASE_ROW_HEIGHT * max(1, n_lines)


def compute_product_row_height(ws, col, desc_text, note_text, min_height):
    """Tinh chieu cao dong san pham can thiet de hien thi HET mo ta san pham
    (font thuong) + ghi chu tinh trang hang (font nghieng, cach 1 dong trong)
    trong cot `col`, khong bi cat/de len dong ben duoi. Luon >= min_height (de
    du cho hinh anh san pham)."""
    font_size_pt = ws[f"{col}{FIRST_PRODUCT_ROW}"].font.sz or 9
    max_width_px = max(20, _merged_width_px(ws, col) - 10)

    regular_font = _load_font("regular", round(font_size_pt * 96 / 72))
    n_lines = _count_multiline_wrapped(desc_text, regular_font, max_width_px)

    if note_text:
        italic_font = _load_font("italic", round(font_size_pt * 96 / 72))
        n_lines += 1 + _count_multiline_wrapped(note_text, italic_font, max_width_px)  # +1 = dong trong cach doan

    line_height_pt = font_size_pt * 1.36  # ty le da kiem chung voi BASE_ROW_HEIGHT (15pt / 11pt)
    needed = max(1, n_lines) * line_height_pt + 10  # +10 le tren/duoi o
    return max(min_height, needed)


def standardize_availability_note(raw):
    """Chuan hoa ghi chu tinh trang hang (co san / dat san xuat) tu text tho
    nguoi dung nhap (vd tren Telegram: "co san", "6-9 ngay") thanh 1 cau hoan
    chinh de chen vao cuoi o "CHI TIET SAN PHAM" (cach dong, chu do, in nghieng
    - xem TextBlock o cho goi ham nay). Tra ve chuoi rong neu raw rong."""
    if not raw:
        return ""
    text = str(raw).strip()
    if not text:
        return ""
    if re.search(r"c[óo]\s*s[ẵa]n", text, re.IGNORECASE):
        return "Hàng có sẵn - giao hàng sau 1-2 ngày kể từ thời điểm xác nhận đơn hàng."
    m = re.search(r"(\d+)\s*-\s*(\d+)\s*ng[àa]y", text, re.IGNORECASE)
    if m:
        return f"Hàng đặt sản xuất {m.group(1)}-{m.group(2)} ngày làm việc"
    m2 = re.search(r"(\d+)\s*ng[àa]y", text, re.IGNORECASE)
    if m2:
        return f"Hàng đặt sản xuất {m2.group(1)} ngày làm việc"
    # Khong nhan dang duoc mau "co san" / "X ngay" / "X-Y ngay" -> giu nguyen
    # text nguoi dung nhap, van in do+nghieng de nguoi xem chu y kiem tra lai.
    return text


def die(msg):
    """Chi dung cho loi nghiem trong KHONG THE tiep tuc (vi du: thieu template,
    danh sach san pham rong). KHONG dung cho loi anh thieu/tai anh loi nua."""
    print(f"LOI: {msg}", file=sys.stderr)
    sys.exit(1)


def maybe_download_image(product, tmp_dir):
    """Neu san pham co image_url thi tai ve, tra ve duong dan file cuc bo.
    Neu khong co anh (thieu du lieu tu web) hoac tai loi, tra ve None thay vi
    crash - de van tao duoc bao gia, chi thieu anh o dong do."""
    if product.get("image_path"):
        return product["image_path"]
    url = product.get("image_url")
    if not url:
        print(
            f"CANH BAO: san pham {product.get('code')} khong co image_path lan "
            f"image_url - bo qua anh cho dong nay, van tiep tuc tao bao gia.",
            file=sys.stderr,
        )
        return None
    import requests
    os.makedirs(tmp_dir, exist_ok=True)
    fname = os.path.join(tmp_dir, f"{product['code'].replace(' ', '_').replace('/', '_')}.png")
    try:
        # MOI: nhieu trang (vd hoaphatsaigon.com) chan hotlink - tu choi/tra ve
        # trang loi (HTML) thay vi anh that neu request khong co header giong
        # trinh duyet that (Referer tu chinh trang do, User-Agent). Neu khong co
        # 2 header nay, requests.get() van co the tra ve HTTP 200 nhung noi dung
        # KHONG PHAI anh -> luu file .png "gia" gay loi crash khi Excel/PIL mo no.
        headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
            ),
            "Referer": "https://hoaphatsaigon.com/",
        }
        resp = requests.get(url, timeout=20, headers=headers)
        resp.raise_for_status()

        # MOI: kiem tra noi dung tai ve CO THUC SU LA ANH HOP LE khong (dung PIL
        # doc thu, khong dua vao Content-Type header hay duoi file - vi co the bi
        # sai/thieu). Neu khong phai anh that (vd bi tra ve trang HTML loi do chan
        # hotlink, link hong, redirect sang trang khac...), bo qua anh cho dong
        # nay thay vi luu file hong roi lam crash ca API o buoc chen anh vao Excel.
        from PIL import Image as PILImage
        import io
        try:
            with PILImage.open(io.BytesIO(resp.content)) as im:
                im.verify()
        except Exception:
            print(
                f"CANH BAO: noi dung tai ve tu {url} khong phai anh hop le "
                f"(co the bi chan hotlink hoac link hong) - bo qua anh cho san pham "
                f"{product.get('code')}, van tiep tuc tao bao gia.",
                file=sys.stderr,
            )
            return None

        with open(fname, "wb") as f:
            f.write(resp.content)
        return fname
    except Exception as e:
        print(
            f"CANH BAO: khong tai duoc anh cho san pham {product.get('code')} "
            f"tu {url} ({e}) - bo qua anh cho dong nay, van tiep tuc tao bao gia.",
            file=sys.stderr,
        )
        return None


def add_centered_image(ws, img_path, row, col_letter, img_width_px, img_height_px, row_height_pts):
    """Chen anh vao giua o (ca chieu ngang va chieu doc), thay vi mac dinh
    openpyxl neo anh o goc tren-trai cua o. Tinh do lech (offset) dua tren do
    rong cot (doi tu don vi "ky tu" cua Excel sang pixel) va chieu cao dong."""
    img = XLImage(img_path)
    img.width = img_width_px
    img.height = img_height_px

    col_dim = ws.column_dimensions.get(col_letter)
    col_width_chars = col_dim.width if col_dim and col_dim.width else 22.43
    col_width_px = col_width_chars * 7 + 5
    row_height_px = row_height_pts * 96 / 72

    off_x_px = max(0, (col_width_px - img_width_px) / 2)
    off_y_px = max(0, (row_height_px - img_height_px) / 2)

    col_idx0 = column_index_from_string(col_letter) - 1  # openpyxl AnchorMarker: 0-indexed
    row_idx0 = row - 1

    marker = AnchorMarker(
        col=col_idx0, colOff=pixels_to_EMU(off_x_px),
        row=row_idx0, rowOff=pixels_to_EMU(off_y_px),
    )
    size = XDRPositiveSize2D(pixels_to_EMU(img_width_px), pixels_to_EMU(img_height_px))
    img.anchor = OneCellAnchor(_from=marker, ext=size)
    ws.add_image(img)


def unmerge_all(ws):
    ranges = [str(mc) for mc in ws.merged_cells.ranges]
    for r in ranges:
        ws.unmerge_cells(r)
    return ranges


def new_row_after_splice(old_row, threshold_row, shift):
    """Tinh vi tri dong moi sau khi chen (shift > 0) hoac xoa (shift < 0) dong
    tai threshold_row. Voi truong hop xoa, vung bi xoa la
    [threshold_row + shift, threshold_row - 1] (shift la so am) - tra ve None
    neu dong do nam trong vung bi xoa (khong con ton tai nua)."""
    if shift >= 0:
        return old_row + shift if old_row >= threshold_row else old_row
    delete_start = threshold_row + shift
    if old_row < delete_start:
        return old_row
    if old_row < threshold_row:
        return None
    return old_row + shift


def remerge_shifted(ws, old_ranges, threshold_row, shift):
    """shift duong: chen them dong. shift am: xoa bot dong."""
    for r in old_ranges:
        min_col, min_row, max_col, max_row = range_boundaries(r)
        new_min_row = new_row_after_splice(min_row, threshold_row, shift)
        new_max_row = new_row_after_splice(max_row, threshold_row, shift)
        if new_min_row is None or new_max_row is None:
            continue  # dong bi xoa, khong re-merge nua
        new_range = f"{get_column_letter(min_col)}{new_min_row}:{get_column_letter(max_col)}{new_max_row}"
        ws.merge_cells(new_range)


def shift_row_heights(ws, old_heights, threshold_row, shift):
    """openpyxl KHONG tu dong dich chuyen row_dimensions (chieu cao dong) khi
    chen/xoa dong bang insert_rows/delete_rows - chi dich chuyen gia tri o va
    KHONG dich merged cells (da xu ly rieng o remerge_shifted). Ham nay dich
    chuyen chieu cao dong tuong ung, tranh tinh trang dong "Tong cong"/"Chiet
    khau" bi ke thua nham chieu cao cao cua dong san pham (anh) sau khi xoa."""
    new_heights = {}
    for old_row, h in old_heights.items():
        new_row = new_row_after_splice(old_row, threshold_row, shift)
        if new_row is None:
            continue  # dong nay da bi xoa, bo qua
        new_heights[new_row] = h

    # Xoa het custom height hien tai (co the con sot lai gia tri cu sai vi tri)
    for row_idx in list(ws.row_dimensions.keys()):
        ws.row_dimensions[row_idx].height = None

    for row_idx, h in new_heights.items():
        ws.row_dimensions[row_idx].height = h


def build(config, tmp_dir="/tmp/bao_gia_images"):
    template_path = config["template_path"]
    output_path = config["output_path"]
    products = config["products"]
    n_products = len(products)

    if n_products == 0:
        die("Danh sach san pham rong - khong the tao bao gia")

    wb = openpyxl.load_workbook(template_path, keep_vba=True)
    ws = wb[SHEET_NAME]

    # --- Thong tin khach hang ---
    phone = config.get("phone", "")
    customer_name = config.get("customer_name", "")
    if customer_name:
        name_line = f"Kính gửi : {customer_name}"
        if phone:
            name_line += f" - {phone}"
    else:
        # Khong co ten khach (vd khach chi dua MST + dia chi + san pham) -> de trong
        # phan sau nhan, KHONG ghep so dien thoai vao day (so dien thoai da co rieng
        # o dong "Dien thoai:" ben duoi).
        name_line = "Kính gửi :"
    ws["A10"] = name_line
    set_wrapped_row_height(ws, 10, name_line)
    if config.get("company"):
        company_line = f"Công Ty: {config['company']}"
        ws["A11"] = company_line
        set_wrapped_row_height(ws, 11, company_line)
    if config.get("tax_address"):
        tax_line = f"Địa chỉ thuế: {config['tax_address']}"
        ws["A12"] = tax_line
        set_wrapped_row_height(ws, 12, tax_line)
    delivery_line = f"Địa chỉ giao hàng: {config.get('delivery_address', '')}"
    ws["A13"] = delivery_line
    set_wrapped_row_height(ws, 13, delivery_line)
    if config.get("mst"):
        mst_line = f"MST: {config['mst']}"
        ws["A14"] = mst_line
        set_wrapped_row_height(ws, 14, mst_line)
    email = config.get("email", "")
    ws["A15"] = f"Điện thoại: {phone}" + " " * 60 + f"Email: {email}"

    boundary_row = FIRST_PRODUCT_ROW + DEFAULT_TEMPLATE_ROWS  # = 21, ranh gioi goc

    # --- Neu so san pham khac 3 (mac dinh cua template), chen hoac XOA HAN dong
    # de khong con dong san pham rong hien thi tren PDF ---
    if n_products != DEFAULT_TEMPLATE_ROWS:
        old_ranges = unmerge_all(ws)
        old_heights = {
            r: dim.height for r, dim in ws.row_dimensions.items() if dim.height is not None
        }
        if n_products > DEFAULT_TEMPLATE_ROWS:
            n_diff = n_products - DEFAULT_TEMPLATE_ROWS
            ws.insert_rows(boundary_row, n_diff)
            shift = n_diff
            # copy dinh dang tu dong san pham cuoi cung co san sang cac dong moi
            last_existing_row = FIRST_PRODUCT_ROW + DEFAULT_TEMPLATE_ROWS - 1  # = 20
            for new_row in range(boundary_row, boundary_row + n_diff):
                ws.row_dimensions[new_row].height = ws.row_dimensions[last_existing_row].height
                for c in range(1, 10):
                    src = ws.cell(row=last_existing_row, column=c)
                    dst = ws.cell(row=new_row, column=c)
                    dst._style = copy(src._style)
        else:
            # It san pham hon template -> XOA HAN cac dong du (khong chi xoa gia tri),
            # de PDF khong con hien o vien trong.
            n_diff = DEFAULT_TEMPLATE_ROWS - n_products
            delete_at = FIRST_PRODUCT_ROW + n_products
            ws.delete_rows(delete_at, n_diff)
            shift = -n_diff
        remerge_shifted(ws, old_ranges, boundary_row, shift)
        shift_row_heights(ws, old_heights, boundary_row, shift)

    # --- Dien du lieu san pham ---
    for idx, product in enumerate(products, start=1):
        row = FIRST_PRODUCT_ROW + idx - 1
        ws[f"A{row}"] = idx
        ws[f"B{row}"] = product["code"]
        desc_text = product.get("description", "")
        note_text = standardize_availability_note(product.get("availability_note", ""))
        if note_text:
            # Cach 1 dong trong, chu do + in nghieng cho phan ghi chu tinh
            # trang hang (co san / dat san xuat) - dung rich text vi phan mo
            # ta san pham van giu font/mau binh thuong.
            note_font = InlineFont(i=True, color="FFFF0000")
            ws[f"D{row}"] = CellRichText(desc_text, "\n\n", TextBlock(note_font, note_text))
        else:
            ws[f"D{row}"] = desc_text
        ws[f"E{row}"] = product.get("color", "")
        ws[f"F{row}"] = "cái"
        ws[f"G{row}"] = product["qty"]
        ws[f"H{row}"] = product["unit_price"]
        ws[f"I{row}"] = f"=H{row}*G{row}"

        # Chieu cao dong PHAI du de hien het mo ta + ghi chu tinh trang hang
        # (khong chi co dinh 170 nhu truoc - vi mo ta dai + co ghi chu se bi
        # cat/de len dong "Tong cong" ben duoi neu chi dung 170 co dinh).
        row_height = compute_product_row_height(ws, "D", desc_text, note_text, ROW_HEIGHT_WITH_IMAGE)
        ws.row_dimensions[row].height = row_height

        img_path = maybe_download_image(product, tmp_dir)
        if img_path:
            add_centered_image(ws, img_path, row, "C", IMG_WIDTH, IMG_HEIGHT, row_height)

    # --- Cong thuc khu vuc tong ket - LUON tinh dong theo vi tri thuc te sau khi
    # chen/xoa dong, khong con nhanh dac biet nua (vi gio luon xoa hoac chen dong
    # cho khop dung n_products, khong con truong hop "dong rong con lai" nua) ---
    shipping_fee = config.get("shipping_fee", 0)
    discount_percent = config.get("discount_percent", 0) or 0
    total_start = FIRST_PRODUCT_ROW + n_products  # dong "Tong cong"
    last_product_row = FIRST_PRODUCT_ROW + n_products - 1
    r_tong_cong = total_start
    r_chiet_khau = total_start + 1
    r_gia_tri_con_lai = total_start + 2
    r_phi_vc = total_start + 3
    r_tong_cong_vc = total_start + 4
    r_vat = total_start + 5
    r_tong_tt_cuoi = total_start + 6
    r_tam_ung = total_start + 7
    r_con_lai = total_start + 8

    ws[f"I{r_tong_cong}"] = f"=SUM(I{FIRST_PRODUCT_ROW}:I{last_product_row})"
    # Chiet khau: neu co discount_percent thi dien cong thuc =Tong_cong*X%
    # (tu dong tinh lai theo tong cong, giong cach lam voi VAT/tam ung o duoi)
    # va HIEN 2 dong nay; neu khong co chiet khau thi AN (hidden) han 2 dong
    # "Chiet khau" va "Gia tri con lai" theo yeu cau - cong thuc ben duoi van
    # giu nguyen (Gia tri con lai = Tong cong - 0 = Tong cong) nen khong pha
    # cac cong thuc VAT/tam ung phia sau, chi la khong hien thi tren PDF.
    if discount_percent:
        ws[f"I{r_chiet_khau}"] = f"=I{r_tong_cong}*{discount_percent}%"
        ws.row_dimensions[r_chiet_khau].hidden = False
        ws.row_dimensions[r_gia_tri_con_lai].hidden = False
    else:
        ws[f"I{r_chiet_khau}"] = 0
        ws.row_dimensions[r_chiet_khau].hidden = True
        ws.row_dimensions[r_gia_tri_con_lai].hidden = True
    ws[f"I{r_gia_tri_con_lai}"] = f"=I{r_tong_cong}-I{r_chiet_khau}"
    ws[f"I{r_phi_vc}"] = shipping_fee
    ws[f"I{r_tong_cong_vc}"] = f"=I{r_gia_tri_con_lai}+I{r_phi_vc}"
    ws[f"I{r_vat}"] = f"=I{r_tong_cong_vc}*8%"
    ws[f"I{r_tong_tt_cuoi}"] = f"=I{r_tong_cong_vc}+I{r_vat}"
    ws[f"I{r_tam_ung}"] = f"=I{r_tong_tt_cuoi}*50%"
    ws[f"I{r_con_lai}"] = f"=I{r_tong_tt_cuoi}-I{r_tam_ung}"

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    wb.save(output_path)

    # --- Kiem tra khong lam mat anh goc cua template ---
    n_media = len([n for n in zipfile.ZipFile(output_path).namelist() if "media" in n])
    print(f"Da luu: {output_path}")
    print(f"So luong anh (media) trong file: {n_media} (phai >= 29 + so anh san pham tai duoc, neu template goc co 29 anh trang tri)")
    if n_media < 29:
        print("CANH BAO: so luong anh thap bat thuong - co the da mat anh goc cua template!", file=sys.stderr)

    return output_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", required=True, help="Duong dan file JSON config")
    ap.add_argument("--tmp-dir", default="/tmp/bao_gia_images", help="Thu muc tam de tai anh san pham")
    args = ap.parse_args()

    with open(args.config, "r", encoding="utf-8") as f:
        config = json.load(f)

    build(config, tmp_dir=args.tmp_dir)


if __name__ == "__main__":
    main()
